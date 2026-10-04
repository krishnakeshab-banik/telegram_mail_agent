"""Extract text from allowed attachment types."""

import csv
from pathlib import Path


def extract_text(path: Path, mime_type: str) -> str:
    """Extract plain text from a file that already passed the allowlist.

    Args:
        path: Temporary file path.
        mime_type: IANA media type.

    Returns:
        Extracted text. Empty when the file has no text layer.
    """
    if mime_type == "application/pdf":
        return _pdf(path)
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _docx(path)
    if mime_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        return _xlsx(path)
    if mime_type in {"text/plain", "text/csv"}:
        return _text(path, mime_type)
    return ""


def _pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _docx(path: Path) -> str:
    from docx import Document

    document = Document(str(path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                cells = [str(cell) for cell in row if cell is not None]
                if cells:
                    lines.append(", ".join(cells))
    finally:
        workbook.close()
    return "\n".join(lines)


def _text(path: Path, mime_type: str) -> str:
    raw = path.read_text(encoding="utf-8", errors="replace")
    if mime_type != "text/csv":
        return raw
    rows = [", ".join(row) for row in csv.reader(raw.splitlines())]
    return "\n".join(rows)
